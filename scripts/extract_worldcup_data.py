from __future__ import annotations

import datetime as dt
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = Path(r"C:\Users\p.struski\Downloads\WorldCup_2026.xlsx")
DEFAULT_OUTPUT = ROOT / "data" / "worldcup-2026.json"


def excel_datetime(value):
    if isinstance(value, dt.datetime):
        return value.strftime("%Y-%m-%dT%H:%M:%S")
    if isinstance(value, dt.date):
        return dt.datetime.combine(value, dt.time()).strftime("%Y-%m-%dT%H:%M:%S")
    return None


def clean_players(players):
    cleaned = []
    seen = set()
    for player in players:
        if not isinstance(player, str):
            continue
        name = player.strip()
        if not name or name.lower() == "nazwa" or name in seen:
            continue
        seen.add(name)
        cleaned.append(name)
    return cleaned


def extract_groups(workbook):
    sheet = workbook["Groups"]
    groups = {}
    current = None
    for row in sheet.iter_rows(values_only=True):
        slot, team_no, name = row[1], row[2], row[3]
        if isinstance(name, str) and re.fullmatch(r"[A-L]", name.strip()):
            current = name.strip()
            groups[current] = []
            continue
        if current and isinstance(slot, str) and isinstance(name, str) and re.fullmatch(r"[A-L][1-4]", slot.strip()):
            groups[current].append(
                {
                    "slot": slot.strip(),
                    "teamNo": int(team_no) if isinstance(team_no, (int, float)) else None,
                    "name": name.strip(),
                }
            )
    return groups


def extract_matches(workbook, groups):
    stage_map = {
        "Round of 32": "1/16 finału",
        "Round of 16": "1/8 finału",
        "1/4 finals": "Ćwierćfinały",
        "1/2 finals": "Półfinały",
        "Third Place": "Mecz o 3. miejsce",
        "Final": "Finał",
    }
    slot_to_team = {team["slot"]: team["name"] for items in groups.values() for team in items}
    sheet = workbook["Matches"]
    stage = "Faza grupowa"
    matches = []
    for row in sheet.iter_rows(values_only=True):
        marker = row[1]
        if isinstance(marker, str) and marker in stage_map:
            stage = stage_map[marker]
            continue
        if not isinstance(marker, (int, float)):
            continue

        home_slot = str(row[2]).strip() if row[2] is not None else ""
        away_slot = str(row[3]).strip() if row[3] is not None else ""
        group = home_slot[0] if re.fullmatch(r"[A-L][1-4]", home_slot) else None
        matches.append(
            {
                "id": int(marker),
                "stage": stage,
                "group": group,
                "homeSlot": home_slot,
                "awaySlot": away_slot,
                "homeTeam": str(row[8]).strip() if row[8] else slot_to_team.get(home_slot, home_slot),
                "awayTeam": str(row[9]).strip() if row[9] else slot_to_team.get(away_slot, away_slot),
                "hostTime": excel_datetime(row[4]),
                "polandTime": excel_datetime(row[5]),
                "venueNo": int(row[6]) if isinstance(row[6], (int, float)) else None,
                "venue": str(row[7]).strip() if row[7] else "",
                "order": row[10] if row[10] is not None else None,
            }
        )
    return matches


def extract_scoring(workbook):
    sheet = workbook["Typer Ustawienia"]
    labels = {
        "Wygrana/Remis/Przegrana:": "outcome",
        "różnica bramek:": "goalDifference",
        "dokładne gole:": "exactGoals",
        "Wiele celów – dobre przybliżenie": "manyGoals",
        "prawidłowa drużyna (faza pucharowa):": "knockoutTeam",
        "Zwycięzca finału/trzeciego miejsca:": "finalWinner",
    }
    scoring = {
        "outcome": 5,
        "goalDifference": 5,
        "exactGoals": 10,
        "manyGoals": 3,
        "knockoutTeam": 10,
        "finalWinner": 10,
    }
    thresholds = {"drawGoals": 4, "bigGoalDifference": 4, "manyGoals": 8}
    for row in sheet.iter_rows(values_only=True):
        label = row[1]
        if isinstance(label, str) and label in labels and isinstance(row[5], (int, float)):
            scoring[labels[label]] = int(row[5])
        if label == "Minimalna liczba bramek w przypadku remisu:" and isinstance(row[12], (int, float)):
            thresholds["drawGoals"] = int(row[12])
        if label == "Minimalna liczba dla dużej różnicy bramek:" and isinstance(row[12], (int, float)):
            thresholds["bigGoalDifference"] = int(row[12])
        if label == "Minimalna liczba dla dużej sumy celów:" and isinstance(row[12], (int, float)):
            thresholds["manyGoals"] = int(row[12])
    return scoring, thresholds


def extract_players(workbook):
    players = []
    for sheet_name in ("Typer_Ranking_1", "Typer_Ranking_2"):
        sheet = workbook[sheet_name]
        for row in sheet.iter_rows(min_row=3, values_only=True):
            players.append(row[2])
    return clean_players(players)


def main():
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SOURCE
    output = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_OUTPUT
    workbook = load_workbook(source, data_only=True, read_only=True)
    groups = extract_groups(workbook)
    scoring, thresholds = extract_scoring(workbook)
    payload = {
        "meta": {
            "name": "Mundial 2026",
            "sourceWorkbook": source.name,
            "generatedAt": dt.datetime.now().isoformat(timespec="seconds"),
            "timezone": "Europe/Warsaw",
            "notes": "Dane bazowe wyeksportowane z arkusza WorldCup_2026.xlsx.",
        },
        "groups": groups,
        "matches": extract_matches(workbook, groups),
        "scoring": scoring,
        "thresholds": thresholds,
        "defaultPlayers": extract_players(workbook),
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Saved {output} ({len(payload['matches'])} matches, {len(payload['defaultPlayers'])} players)")


if __name__ == "__main__":
    main()
